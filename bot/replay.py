"""Replay / backtest: feed historical 1m data through the identical signal and paper-execution
code path and write the same JSON schema into a separate directory (default ``data_replay/``).

Input files (``.xlsx`` or ``.csv``) have the columns
``time, THETA_open, THETA_high, THETA_low, THETA_close, THETA_volume, TFUEL_open, ...,
TFUEL_volume, THETA_per_TFUEL``. ``THETA_per_TFUEL`` is ignored; the ratio is recomputed from
the closes. An ``.xlsx`` file may hold several sheets (Excel's 1,048,576-row limit); all are read.

There is no order book in historical data, so ``mid`` and ``touch`` both fill at the bar close
(the "close fill" / optimistic case) and ``worst`` fills at the next bar's high/low.
"""

from __future__ import annotations

import csv
import itertools
import logging
import time
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

from bot.config import Settings
from bot.instance import Instance
from bot.market_data import JointBar, Kline, bar_fields
from bot.portfolio import Portfolio, prices, value_of
from bot.reference import reference_range
from bot.signals import MarketEma
from bot.storage import Storage, atomic_write_json
from bot.util import SimClock, iso_to_ms, ms_to_iso, year_of, year_start_ms

log = logging.getLogger("bot.replay")

COLUMNS = [f"{a}_{f}" for a in ("THETA", "TFUEL") for f in ("open", "high", "low", "close", "volume")]
EXCEL_EPOCH = datetime(1899, 12, 30, tzinfo=timezone.utc)


def parse_time(v: Any) -> int:
    """Open time of the minute as epoch ms (UTC)."""
    if isinstance(v, datetime):
        dt = v if v.tzinfo else v.replace(tzinfo=timezone.utc)
        return int(round(dt.timestamp() * 1000))
    if isinstance(v, (int, float)):
        x = float(v)
        if x > 1e14:  # microseconds
            return int(x / 1000)
        if x > 1e11:  # milliseconds
            return int(x)
        if x > 1e8:  # seconds
            return int(x * 1000)
        return int(round((EXCEL_EPOCH + timedelta(days=x)).timestamp() * 1000))  # Excel serial date
    s = str(v).strip()
    if s.replace(".", "", 1).isdigit():
        return parse_time(float(s))
    return iso_to_ms(s.replace(" ", "T") if "T" not in s else s)


def _row_to_bar(ts: int, vals: list[float]) -> JointBar:
    th = Kline(ts, vals[0], vals[1], vals[2], vals[3], vals[4])
    tf = Kline(ts, vals[5], vals[6], vals[7], vals[8], vals[9])
    return JointBar(ts, th, tf)


def _header_index(header: list[Any], path: Path) -> tuple[int, list[int]]:
    names = [str(h).strip() if h is not None else "" for h in header]
    lower = [n.lower() for n in names]
    try:
        t_idx = lower.index("time")
    except ValueError:
        t_idx = next((i for i, n in enumerate(lower) if n in ("open_time", "timestamp", "datetime", "date")), -1)
        if t_idx < 0:
            raise ValueError(f"{path}: no 'time' column in header {names}")
    idx = []
    for c in COLUMNS:
        if c.lower() not in lower:
            raise ValueError(f"{path}: missing column {c}")
        idx.append(lower.index(c.lower()))
    return t_idx, idx


def read_csv(path: Path) -> Iterator[JointBar]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        header = next(reader)
        t_idx, idx = _header_index(header, path)
        for row in reader:
            if not row or not row[t_idx].strip():
                continue
            try:
                yield _row_to_bar(parse_time(row[t_idx]), [float(row[i]) for i in idx])
            except (ValueError, IndexError):
                continue


def read_xlsx(path: Path) -> Iterator[JointBar]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = None
            for row in rows:
                if row and any(c is not None for c in row):
                    header = list(row)
                    break
            if header is None:
                continue
            t_idx, idx = _header_index(header, path)
            for row in rows:
                if not row or row[t_idx] is None:
                    continue
                try:
                    yield _row_to_bar(parse_time(row[t_idx]), [float(row[i]) for i in idx])
                except (ValueError, TypeError, IndexError):
                    continue
    finally:
        wb.close()


def _reader(p: Path) -> Iterator[JointBar]:
    return read_xlsx(p) if p.suffix.lower() in (".xlsx", ".xlsm") else read_csv(p)


def _first_ts(p: Path) -> int:
    for b in _reader(p):
        return b.ts_ms
    return 2**62


def iter_bars(paths: Iterable[Path | str], stats: dict[str, int] | None = None) -> Iterator[JointBar]:
    """Stream bars from all files in time order without loading everything into memory.

    Files are ordered by their first timestamp; rows are expected in ascending time order
    within a file (as exported). Duplicate or out-of-order rows are skipped and counted.
    """
    stats = stats if stats is not None else {}
    stats.setdefault("rows", 0)
    stats.setdefault("skipped", 0)
    files = sorted((Path(p) for p in paths), key=_first_ts)
    last = -1
    for p in files:
        t0 = time.time()
        n = 0
        for b in _reader(p):
            if b.theta.c <= 0 or b.tfuel.c <= 0 or b.ts_ms <= last:
                stats["skipped"] += 1
                continue
            last = b.ts_ms
            n += 1
            stats["rows"] += 1
            yield b
        log.info("read %d rows from %s in %.1fs", n, p, time.time() - t0)


def load_bars(paths: Iterable[Path | str]) -> list[JointBar]:
    return list(iter_bars(paths))


class YearTracker:
    """Per-year excess for each portfolio variant.

    * ``excess_rebased``: HODL benchmark reset to the holdings at the start of the year, i.e.
      "trading this year vs. simply holding what we had on 1 January".
    * ``excess_chained``: growth of (V/V_hodl) over the year with the original benchmark.
    """

    def __init__(self) -> None:
        self.open: dict[tuple[str, str], dict[str, Any]] = {}
        self.results: list[dict[str, Any]] = []

    def start(self, inst: str, variant: str, port: Portfolio, px: dict[str, float], ts: int) -> None:
        self.open[(inst, variant)] = {
            "year": year_of(ts),
            "start_ts": ts,
            "start_bal": dict(port.bal),
            "start_ratio": 1.0 + port.excess(px),
            "start_rebalances": port.rebalances_total,
            "start_fees": port.fees_usd,
        }

    def close(self, inst: str, variant: str, port: Portfolio, px: dict[str, float], ts: int) -> None:
        o = self.open.pop((inst, variant), None)
        if not o:
            return
        v = port.value(px)
        hv_y = value_of(o["start_bal"], px)
        self.results.append(
            {
                "instance": inst,
                "variant": variant,
                "year": o["year"],
                "start": ms_to_iso(o["start_ts"]),
                "end": ms_to_iso(ts),
                "days": round((ts - o["start_ts"]) / 86_400_000, 2),
                "excess_rebased": v / hv_y - 1.0 if hv_y > 0 else None,
                "excess_chained": (1.0 + port.excess(px)) / o["start_ratio"] - 1.0,
                "trades": port.rebalances_total - o["start_rebalances"],
                "fees_usd": port.fees_usd - o["start_fees"],
            }
        )


def run_replay(
    settings: Settings,
    files: list[str | Path],
    instance_names: list[str],
    out_dir: str | Path = "data_replay",
    start: str | None = None,
    end: str | None = None,
    overlay_warmup_bars: int | None = None,
    ladder_warmup_days: int | None = None,
    progress: bool = True,
) -> dict[str, Any]:
    out = Path(out_dir)
    if out.exists() and any(out.iterdir()):
        raise SystemExit(f"replay output directory {out} is not empty; remove it or pass --out <new dir>")
    t0 = time.time()
    stats: dict[str, int] = {}
    bars = iter_bars(files, stats)
    first = next(bars, None)
    if first is None:
        raise SystemExit("no rows read from input files")
    start_ms = iso_to_ms(start) if start else None
    end_ms = iso_to_ms(end) if end else None

    clock = SimClock(first.ts_ms)
    storage = Storage(out, clock, flush_each=False)
    storage.log_events = False
    storage.event("INFO", "replay_start", f"replay of {[str(f) for f in files]}",
                  ts_ms=first.ts_ms, instances=instance_names)

    instances: list[Instance] = []
    for name in instance_names:
        cfg = settings.instance(name)
        s = cfg.strategy
        if overlay_warmup_bars is not None:
            s = replace(s, overlay_warmup_bars=overlay_warmup_bars)
        if ladder_warmup_days is not None:
            s = replace(s, ladder_warmup_days=ladder_warmup_days)
        cfg = replace(cfg, strategy=s, live_requested=False, api_key="", api_secret="")
        instances.append(Instance(cfg, storage, clock, replay=True, summary_every_min=1440, state_every_min=10**12))

    market_ema = MarketEma(settings.strategy.overlay_ema_span_min)
    years = YearTracker()
    next_year_ms: dict[str, int] = {}
    last_bar = first
    last_px = prices(first.theta.c, first.tfuel.c)
    n_bars = 0

    for bar in itertools.chain([first], bars):
        if end_ms is not None and bar.ts_ms > end_ms:
            break
        clock.set(bar.close_ms)
        px = prices(bar.theta.c, bar.tfuel.c)
        ema, dev = market_ema.update(bar.ts_ms, bar.lr)
        storage.write_bar(bar.ts_ms, **bar_fields(bar, ema, dev))
        phase = "warmup" if start_ms is not None and bar.ts_ms < start_ms else "live"
        for inst in instances:
            was_started = inst.started
            inst.on_bar(bar, phase, px)
            if not inst.started:
                continue
            if not was_started:
                for v, p in inst.paper.items():
                    years.start(inst.name, v, p, px, bar.close_ms)
                next_year_ms[inst.name] = year_start_ms(year_of(bar.ts_ms) + 1)
            elif bar.ts_ms >= next_year_ms[inst.name]:  # first bar of a new UTC year
                # Year boundary: close with the previous bar's prices, open the new year.
                for v, p in inst.paper.items():
                    years.close(inst.name, v, p, last_px, last_bar.close_ms)
                    years.start(inst.name, v, p, last_px, last_bar.close_ms)
                next_year_ms[inst.name] = year_start_ms(year_of(bar.ts_ms) + 1)
        last_bar, last_px = bar, px
        n_bars += 1
        if progress and n_bars % 100_000 == 0:
            log.info("replay: %d bars, at %s", n_bars, ms_to_iso(bar.ts_ms))

    for inst in instances:
        if inst.started:
            for v, p in inst.paper.items():
                years.close(inst.name, v, p, last_px, last_bar.close_ms)
        inst.write_summary(last_bar.close_ms, last_px)
        inst.save()
    storage.event("INFO", "replay_end", "replay finished", ts_ms=last_bar.close_ms)
    storage.save_runtime(
        {
            "ts": ms_to_iso(last_bar.close_ms),
            "schema_version": 1,
            "replay": True,
            "ws": {"ws_connected": None, "lag_sec": None, "last_bar_ts": ms_to_iso(last_bar.ts_ms)},
            "kill_switch": {"active": False},
            "instances": {i.name: i.mode for i in instances},
        }
    )
    storage.close()

    report = {
        "generated": ms_to_iso(int(time.time() * 1000)),
        "files": [str(f) for f in files],
        "bars": n_bars,
        "rows_skipped": stats.get("skipped", 0),
        "first_bar": ms_to_iso(first.ts_ms),
        "last_bar": ms_to_iso(last_bar.ts_ms),
        "trading_window": {"start": start, "end": end},
        "instances": {},
        "years": years.results,
        "notes": [
            "mid and touch both fill at the bar close (no historical order book) = optimistic case",
            "worst fills at the next bar's high (buys) / low (sells) = pessimistic case",
            "excess_rebased: HODL reset to holdings on 1 Jan; excess_chained: growth of V/V_hodl over the year",
            "the backtest reference also includes a market-impact model that this replay does not reproduce "
            "(set PAPER_IMPACT_BPS_PER_1K to approximate one)",
        ],
        "elapsed_sec": round(time.time() - t0, 1),
    }
    for inst in instances:
        st = inst.engine
        report["instances"][inst.name] = {
            "capital_usd": inst.cfg.capital_usd,
            "trading_started": ms_to_iso(inst.trading_started_ms),
            "overlay_warmup_bars": inst.cfg.strategy.effective_overlay_warmup_bars,
            "ladder_warmup_days": inst.cfg.strategy.ladder_warmup_days,
            "excess_full_period": {v: p.excess(last_px) for v, p in inst.paper.items()},
            "trades": {v: p.rebalances_total for v, p in inst.paper.items()},
            "fees_usd": {v: p.fees_usd for v, p in inst.paper.items()},
            "final_ladder_w": st.ladder_w,
            "final_pos": st.overlay.pos,
        }
    atomic_write_json(out / "replay_report.json", report, indent=2)
    return report



def format_report(report: dict[str, Any]) -> str:
    lines = [
        f"Replay {report['first_bar']} .. {report['last_bar']}  ({report['bars']} bars, {report['elapsed_sec']}s)",
        "",
        f"{'inst':5} {'year':5} {'days':>6} {'worst':>9} {'close':>9} {'reference':>14}  {'trades':>6}   (excess vs HODL, rebased per year)",
    ]
    by = {(r["instance"], r["year"], r["variant"]): r for r in report["years"]}
    keys = sorted({(r["instance"], r["year"]) for r in report["years"]})
    for inst, year in keys:
        w = by.get((inst, year, "worst"))
        m = by.get((inst, year, "mid"))
        ref = reference_range(inst, year)
        ref_s = f"{ref[0]*100:+.0f}..{ref[1]*100:+.0f}%" if ref else "-"
        lines.append(
            f"{inst:5} {year:5} {(w or m or {}).get('days', 0):6.0f} "
            f"{_pct(w and w['excess_rebased']):>9} {_pct(m and m['excess_rebased']):>9} {ref_s:>14}  "
            f"{(m or {}).get('trades', 0):>6}"
        )
    lines.append("")
    for name, info in report["instances"].items():
        ex = ", ".join(f"{v} {_pct(x)}" for v, x in info["excess_full_period"].items())
        lines.append(f"{name}: full period excess: {ex}; trading started {info['trading_started']}")
    return "\n".join(lines)


def _pct(x: float | None) -> str:
    return "-" if x is None else f"{x*100:+.2f}%"
